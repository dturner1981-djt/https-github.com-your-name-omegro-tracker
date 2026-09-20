"""Parse the quarterly Operational Governance Scorecard.

The scorecard is a wide per-quarter matrix: one row band per VBU, metric names
across the top ("Net Revenue", "EBITA %", "Organic Growth %", "Core Ratio",
"WC %", ...), plus the governance stage and score.

Unlike the other parsers this one scans rather than indexes fixed cells. The
scorecard is rebuilt each quarter and columns move, so the parser locates the
header row by looking for a run of known metric labels, then reads each BU's
row by matching column A against the alias list.

NOTE: the layout heuristics here were written against the published metric
vocabulary, not against a downloaded copy of the Q2-26 workbook (it is several
megabytes and was not retrievable at build time). Run
`omegro-tracker validate --source og_scorecard` after the first authenticated
refresh; it prints what the scanner matched so the mapping can be confirmed
before the numbers are trusted.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..model import Fact, Governance
from .monthly_review import PCT_METRICS, _number, _squash

# Scorecard column labels -> internal metric keys.
COLUMN_ALIASES: dict[str, str] = {
    "netrevenue": "net_revenue",
    "nr": "net_revenue",
    "ebita": "ebita",
    "ebita%": "ebita_pct",
    "organicgrowth%": "og_pct",
    "og%": "og_pct",
    "organicgrowth": "og_pct",
    "coreratio": "core_ratio",
    "bqr%": "bqr_pct",
    "bqr": "bqr_pct",
    "wc%": "wc_pct",
    "workingcapital%": "wc_pct",
    "maintenanceratio": "maintenance_ratio",
    "g&a%": "ga_pct",
    "headcount": "headcount",
}

STAGE_LABELS = ("invest", "grow", "monitor", "takeaction", "escalation")

# A header row must carry at least this many recognised metric labels before we
# accept it — the scorecard has several decorative banner rows above the data.
_MIN_HEADER_HITS = 3


def _alias_index(units: list[dict[str, Any]]) -> dict[str, str]:
    index: dict[str, str] = {}
    for unit in units:
        for alias in [unit["name"], unit.get("short_name", ""), *unit.get("aliases", [])]:
            if alias:
                index[re.sub(r"[^a-z0-9]", "", alias.lower())] = unit["key"]
    return index


def _find_header(sheet: Any, scan_rows: int = 40, scan_cols: int = 60) -> tuple[int, dict[int, str]] | None:
    best: tuple[int, dict[int, str]] | None = None
    for r in range(1, scan_rows + 1):
        columns: dict[int, str] = {}
        for c in range(1, scan_cols + 1):
            if key := COLUMN_ALIASES.get(_squash(sheet.cell(row=r, column=c).value)):
                columns[c] = key
        if len(columns) >= _MIN_HEADER_HITS and (best is None or len(columns) > len(best[1])):
            best = (r, columns)
    return best


def _stage(value: Any) -> str | None:
    squashed = _squash(value)
    for stage in STAGE_LABELS:
        if squashed.startswith(stage):
            text = str(value).strip()
            return text
    return None


def parse(
    path: str | Path,
    *,
    units: list[dict[str, Any]],
    period: str,
    source: str = "og_scorecard",
) -> tuple[list[Fact], list[Governance]]:
    """Returns (facts, governance). `period` is a quarter, e.g. "2026-Q2"."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    aliases = _alias_index(units)

    facts: list[Fact] = []
    governance: dict[str, Governance] = {}

    for name in workbook.sheetnames:
        sheet = workbook[name]
        header = _find_header(sheet)
        if header is None:
            continue
        header_row, columns = header

        max_row = min(sheet.max_row or 0, header_row + 400)
        for r in range(header_row + 1, max_row + 1):
            label = sheet.cell(row=r, column=1).value
            bu = aliases.get(re.sub(r"[^a-z0-9]", "", str(label or "").lower()))
            if not bu:
                continue

            for col, metric in columns.items():
                value = _number(sheet.cell(row=r, column=col).value)
                if value is None:
                    continue
                facts.append(
                    Fact(
                        bu=bu,
                        metric=metric,
                        basis="quarter",
                        measure="actual",
                        period=period,
                        value=value,
                        unit="pct" if metric in PCT_METRICS else "kGBP",
                        source=source,
                    )
                )

            record = governance.setdefault(bu, Governance(bu=bu, period=period))
            # Stage and score sit alongside the metric block; find them by
            # scanning the row rather than assuming a column.
            for c in range(1, 60):
                cell = sheet.cell(row=r, column=c).value
                if record.stage is None and (stage := _stage(cell)):
                    record.stage = stage
                head = _squash(sheet.cell(row=header_row, column=c).value)
                if record.score is None and head in {"ogscore", "score", "og%", "governancescore"}:
                    record.score = _number(cell)

    workbook.close()
    return facts, list(governance.values())


def describe(path: str | Path, units: list[dict[str, Any]]) -> str:
    """Human-readable account of what the scanner matched.

    Used by `omegro-tracker validate` so the column mapping can be eyeballed
    against the real workbook before anyone relies on the figures.
    """
    workbook = load_workbook(path, data_only=True, read_only=True)
    aliases = _alias_index(units)
    lines: list[str] = []
    for name in workbook.sheetnames:
        sheet = workbook[name]
        header = _find_header(sheet)
        if header is None:
            lines.append(f"  {name}: no metric header found")
            continue
        header_row, columns = header
        found = [
            str(sheet.cell(row=r, column=1).value)
            for r in range(header_row + 1, min(sheet.max_row or 0, header_row + 400) + 1)
            if aliases.get(re.sub(r"[^a-z0-9]", "", str(sheet.cell(row=r, column=1).value or "").lower()))
        ]
        cols = ", ".join(f"{c}->{m}" for c, m in sorted(columns.items()))
        lines.append(f"  {name}: header row {header_row}; columns [{cols}]; units {found}")
    workbook.close()
    return "\n".join(lines)
